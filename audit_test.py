# -*- coding: utf-8 -*-
"""
审计核对引擎的离线回归测试（不打网、不加载真实 OCR 模型）。

运行方式：        的前提下
    HOME=$PWD/.smoke_home .venv/bin/python audit_test.py
  （无需启动后端、无需联网、无需 PaddleOCR 模型）

覆盖点：
  1) detect_columns  常见列名 → 角色列号识别
  2) extract_addr_parts  收货地址 / 水印地址的地市切分（含直辖市、无“省”、 吉林市 等边界）
  3) extract_invoice_fields  发票字段提取（三种行序）
  4) run_audit 端到端（打桩网上抓取与 OCR）：
     - 全部一致  → 审核通过-补贴未发放
     - 发票号不一致 → 审核未通过
     - 发票无法提取 → 需人工复核
  5) 原文件（内存副本）不被修改
"""
import io
import sys
import time

import audit_engine as ae

PASSED = []


def check(name, cond, detail=""):
    if cond:
        PASSED.append(name)
        print(f"  ✓ {name}")
    else:
        print(f"  ✗ {name}  {detail}")
        sys.exit(1)


def octets(x):  # 内存字节指纹
    return (x, len(x))


# ---------------------------------------------------------------------------
# 1) 列角色识别
# ---------------------------------------------------------------------------
def test_detect_columns():
    print("[1] detect_columns 列角色识别")
    headers = [
        "人员", "姓名", "身份证号", "手机号", "品类编码", "品类状态",
        "销售数量", "销售金额", "补贴金额", "用户实付金额",
        "交易完成或商品退回时间", "收货地区编码", "收货地址",
        "销售方式", "销售企业名称", "统一社会信息代码",
        "商品条码", "商品名称", "品牌名称", "规格型号",
        "能耗等级", "补贴比例", "S/N码", "", "IMEI1码", "IMEI2码",
        "物流公司", "物流单号", "签收时间", "收货人",
        "收货人手机号", "收货日期", "平台订单编号", "平台支付交易编号",
        "平台退款交易编号", "支付方式", "发票号", "开票日期",
        "发票购买方名称", "发票销售方名称", "开票金额",
        "发票链接地址或发票图片", "物流单链接地址",
        "3C外包装照片", "3C商品图片链接", "3C外包装与商品合照",
        "家电外包装SN照片", "家电外包装妥投照片",
        "退货（款）日期", "备注", "审核情况", "审核备注",
    ]
    m = ae.detect_columns(headers)
    want = {
        "invoice_link": 41, "jd_link": 42, "sn_image_link": 46, "wm_image_link": 47,
        "invoice_no": 36, "invoice_date": 37, "invoice_amount": 40,
        "buyer_name": 38, "seller_name": 39, "sn_code": 22, "receive_addr": 12,
    }
    for role, col in want.items():
        got = m[role]["index"] if m.get(role) else None
        check(f"列[{role}] 应={col} 实={got}", got == col, str(m.get(role)))


# ---------------------------------------------------------------------------
# 2) 地址切分
# ---------------------------------------------------------------------------
def test_addr():
    print("[2] extract_addr_parts 地址切分")
    cases = [
        ("江苏省南京市江宁区东山街道江宁路1号", "南京市", "江宁区"),
        ("江苏南京市江宁区东山街道江宁路1号", "南京市", "江宁区"),
        ("吉林省吉林市丰满区吉林大街", "吉林市", "丰满区"),
        ("北京市海淀区中关村大街", "北京市", "海淀区"),
        ("浙江省杭州市西湖区文三路", "杭州市", "西湖区"),
        ("广东省东莞市长安镇乌沙", "东莞市", ""),
        ("上海市浦东新区张江高科技园区", "上海市", "浦东新区"),
    ]
    for src, city, county in cases:
        _, c, k = ae.extract_addr_parts(src)
        check(f"地址[{src[:12]}...] 市={c} 区县={k}", c == city and k == county,
              f"实得 {c!r}/{k!r}")


# ---------------------------------------------------------------------------
# 3) 发票字段提取（三种行序）
# ---------------------------------------------------------------------------
def test_invoice_fields():
    print("[3] extract_invoice_fields 发票字段（三种行序）")
    want = {"invoice_no": "12345678901234567890", "invoice_date": "2026-02-28",
            "invoice_amount": "908.85",
            "buyer_name": "测试用户甲", "seller_name": "示例贸易有限公司"}
    t1 = ("电子发票(普通发票)\n发票号码：12345678901234567890\n开票日期：2026年02月28日\n"
          "购买方信息\n名称：测试用户甲\n统一社会信用代码/纳税人识别号：110101199001011234\n"
          "销售方信息\n名称：示例贸易有限公司\n"
          "统一社会信用代码/纳税人识别号：91310000MA1K00000X\n"
          "价税合计(大写)玖佰零捌圆捌角伍分(小写)¥908.85\n")
    t2 = ("电子发票(普通发票)\n发票号码：12345678901234567890\n开票日期：2026年02月28日\n"
          "购买方信息\n销售方信息\n名称：测试用户甲\n名称：示例贸易有限公司\n"
          "统一社会信用代码/纳税人识别号：110101199001011234\n"
          "统一社会信用代码/纳税人识别号：91310000MA1K00000X\n(小写)¥908.85\n")
    t3 = ("电子发票（普通发票）\n发票号码:12345678901234567890\n开票日期:2026年02月28日\n"
          "名称:测试用户甲\n统一社会信用代码/纳税人识别号:110101199001011234\n"
          "名称:示例贸易有限公司\n"
          "统一社会信用代码/纳税人识别号:91310000MA1K00000X\n价税合计(小写)¥908.85\n")
    for i, t in enumerate((t1, t2, t3), 1):
        f = ae.extract_invoice_fields(t)
        ok = all(f.get(k) == v for k, v in want.items())
        check(f"行序{t} 全字段一致", ok, str(f))


# ---------------------------------------------------------------------------
# 4) run_audit 端到端（打桩）
# ---------------------------------------------------------------------------
PNG_1PX = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00"
    b"\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\x0aIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01"
    b"\x0d\x0a\x2d\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)

HEADERS = [
    "收货地址", "S/N码", "发票号", "开票日期", "开票金额",
    "发票购买方名称", "发票销售方名称",
    "发票链接地址或发票图片", "物流单链接地址",
    "家电外包装SN照片", "家电外包装妥投照片",
]
MAPPING = {
    "header_row": 1,
    "receive_addr": 1, "sn_code": 2,
    "invoice_no": 3, "invoice_date": 4, "invoice_amount": 5,
    "buyer_name": 6, "seller_name": 7,
    "invoice_link": 8, "jd_link": 9, "sn_image_link": 10, "wm_image_link": 11,
}


def build_xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class StubEnv:
    """打桩：fetch_bytes 返回假图；invoice_fields 返回预制字段；_run_ocr 按调用顺序返回文本。"""

    def __init__(self, invoice_fields, ocr_responses):
        self.invoice_fields = invoice_fields
        self.ocr_queue = list(ocr_responses)
        ae.fetch_bytes = lambda url: PNG_1PX
        ae.invoice_fields = lambda raw: (dict(self.invoice_fields), "stub")
        self.stack = []
        import app as _app
        self.orig_run_ocr = _app._run_ocr
        _app._run_ocr = self.fake_run_ocr

    def fake_run_ocr(self, img):
        return self.ocr_queue.pop(0) if self.ocr_queue else {"items": []}

    def restore(self):
        import app as _app
        ae.fetch_bytes = self.orig_fetch
        ae.invoice_fields = self.orig_inv
        _app._run_ocr = self.orig_run_ocr

    def __enter__(self):
        self.orig_fetch = ae.fetch_bytes
        self.orig_inv = ae.invoice_fields
        return self

    def __exit__(self, *a):
        self.restore()


def test_run_audit_pass():
    print("[4] run_audit 端到端：全部一致 → 通过")
    inv = {"invoice_no": "12345678901234567890", "invoice_date": "2026-02-28",
           "invoice_amount": "908.85", "buyer_name": "测试用户甲",
           "seller_name": "示例贸易有限公司"}
    ocr = [
        {"items": [{"text": "SN码 SN1234567890", "score": 0.99, "box": []}]},
        {"items": [{"text": "拍摄地点：江苏省南京市江宁区", "score": 0.99, "box": []}]},
    ]
    row = ["江苏南京市江宁区东山街道江宁路1号", "SN1234567890",
           "12345678901234567890", "2026-02-28", "908.85",
           "测试用户甲", "示例贸易有限公司",
           "https://inv/x.pdf", "https://jd/x", "https://sn/x", "https://wm/x"]
    xlsx = build_xlsx([row])
    before = octets(xlsx)
    with StubEnv(inv, ocr):
        res = ae.run_audit(xlsx, dict(MAPPING))
    check("通过-统计 pass=1", res["counts"].get("pass") == 1, str(res["counts"]))
    check("通过-状态串", res["rows"][0]["status"] == ae.RESULT_PASS, res["rows"][0]["status"])
    check("通过-原文件未动", octets(xlsx) == before)
    # 结果表里 审核情况/审核备注 已写入
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(res["xlsx"]))
    ws = wb.worksheets[0]
    status_val = ws.cell(row=2, column=ws.max_column - 1).value
    remark = ws.cell(row=2, column=ws.max_column).value
    check("通过-结果表已写状态列", status_val == ae.RESULT_PASS, str(status_val))
    check("通过-结果表备注含发票/S/N/水印", "S/N码一致" in remark and "水印市级一致" in remark, remark[:200])


def test_run_audit_fail():
    print("[5] run_audit 端到端：发票号不一致 → 未通过")
    inv = {"invoice_no": "99999999999999999999", "invoice_date": "2026-02-28",
           "invoice_amount": "908.85", "buyer_name": "测试用户甲",
           "seller_name": "示例贸易有限公司"}
    ocr = [
        {"items": [{"text": "SN1234567890", "score": 0.99, "box": []}]},
        {"items": [{"text": "江苏省南京市江宁区", "score": 0.99, "box": []}]},
    ]
    row = ["南京市江宁区东山街道", "SN1234567890",
           "12345678901234567890", "2026-02-28", "908.85",
           "测试用户甲", "示例贸易有限公司",
           "https://inv/x.pdf", "https://jd/x", "https://sn/x", "https://wm/x"]
    with StubEnv(inv, ocr):
        res = ae.run_audit(build_xlsx([row]), dict(MAPPING))
    check("未通过-统计 fail=1", res["counts"].get("fail") == 1, str(res["counts"]))
    check("未通过-状态串", res["rows"][0]["status"] == ae.RESULT_FAIL, res["rows"][0]["status"])
    check("未通过-备注含发票号不一致",
          "发票号不一致" in res["rows"][0]["remark"], res["rows"][0]["remark"])


def test_run_audit_manual():
    print("[6] run_audit 端到端：发票无法提取 → 需人工复核")
    ocr = [
        {"items": [{"text": "SN1234567890", "score": 0.99, "box": []}]},
        {"items": [{"text": "江苏省南京市江宁区", "score": 0.99, "box": []}]},
    ]
    row = ["南京市江宁区东山街道", "SN1234567890",
           "12345678901234567890", "2026-02-28", "908.85",
           "测试用户甲", "示例贸易有限公司",
           "https://inv/x.pdf", "https://jd/x", "https://sn/x", "https://wm/x"]
    with StubEnv({}, ocr):
        res = ae.run_audit(build_xlsx([row]), dict(MAPPING))
    check("人工复核-统计 manual=1", res["counts"].get("manual") == 1, str(res["counts"]))
    check("人工复核-状态串", res["rows"][0]["status"] == ae.RESULT_MANUAL, res["rows"][0]["status"])


def test_run_audit_empty_link_hint():
    print("[7] run_audit：链接列已映射但为空 → 备注带列号定位")
    inv = {"invoice_no": "12345678901234567890", "invoice_date": "2026-02-28",
           "invoice_amount": "908.85", "buyer_name": "测试用户甲",
           "seller_name": "示例贸易有限公司"}
    ocr = [
        {"items": [{"text": "SN码 SN1234567890", "score": 0.99, "box": []}]},
        {"items": [{"text": "拍摄地点：江苏省南京市江宁区", "score": 0.99, "box": []}]},
    ]
    row = ["南京市江宁区东山街道", "SN1234567890",
           "12345678901234567890", "2026-02-28", "908.85",
           "测试用户甲", "示例贸易有限公司",
           "https://inv/x.pdf", "", "https://sn/x", "https://wm/x"]
    # jd_link 映射到第 12 列（该行没有值）→ 应有“已映射第12列”定位提示
    m = dict(MAPPING); m["jd_link"] = 12
    with StubEnv(inv, ocr):
        res = ae.run_audit(build_xlsx([row]), m)
    rem = res["rows"][0]["remark"]
    check("空链接列带列号定位", "已映射第12列" in rem, rem[:200])
    # 未映射任何列时，提示“未映射…下拉框里没选”
    m2 = dict(MAPPING); m2["wm_image_link"] = None
    ocr2 = [
        {"items": [{"text": "SN码 SN1234567890", "score": 0.99, "box": []}]},
    ]
    with StubEnv(inv, ocr2):
        res2 = ae.run_audit(build_xlsx([row]), m2)
    rem2 = res2["rows"][0]["remark"]
    check("未映射链接列提示未映射", "未映射" in rem2 and "下拉框里没选" in rem2, rem2[:200])


if __name__ == "__main__":
    t0 = time.time()
    test_detect_columns()
    test_addr()
    test_invoice_fields()
    test_run_audit_pass()
    test_run_audit_fail()
    test_run_audit_manual()
    test_run_audit_empty_link_hint()
    print(f"\n全部通过 {len(PASSED)} 项断言，耗时 {time.time() - t0:.1f}s")
